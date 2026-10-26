"""Reconcile exactly 82 original IDs with executed evidence; never infer formal acceptance."""
import collections
import copy
import datetime
import hashlib
import json
import pathlib
import re
import subprocess

ROOT = pathlib.Path(__file__).resolve().parents[3]
OUT = ROOT / 'build/staging/close19'
DOC = ROOT / 'docs/anexo14'
PUBLIC = DOC / 'evidencia-cierre19'
BASE = '005573f16'
FULL = 'IMPLEMENTADO Y VERIFICADO INTEGRALMENTE EN STAGING'
PART = 'IMPLEMENTADO CON EVIDENCIA PARCIAL'
WAIT = 'PENDIENTE DE PRUEBA ADICIONAL'
BLOCK = 'BLOQUEADO POR INSUMO O DECISIÓN EXTERNA'


def read(p):
    return json.loads(p.read_text(encoding='utf-8-sig'))


def dump(p, value):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def git(*args):
    return subprocess.check_output(['git', *args], cwd=ROOT, text=True, encoding='utf-8').strip()


def latest(pattern, predicate=lambda r: True):
    matches = [(p, read(p)) for p in OUT.glob(pattern)]
    return max((pair for pair in matches if predicate(pair[1])), key=lambda pair: pair[0].stat().st_mtime)


def cell(value):
    if isinstance(value, list):
        value = '; '.join(map(str, value))
    return str(value).replace('|', '\\|').replace('\n', ' ')


def main():
    baseline = json.loads(git('show', BASE + ':docs/anexo14/TRAZABILIDAD_82_2026-10-25.json'))
    oldmatrix = git('show', BASE + ':docs/anexo14/MATRIZ_ASODELTOL_82.md')
    priorities = {line.split('|')[1].strip(): line.split('|')[3].strip()
                  for line in oldmatrix.splitlines() if line.startswith('| REQ-')}
    winpath, win = latest('browser-*/results.json', lambda r: len(r['engines']) == 3 and all(e['hostOS'] == 'win32' for e in r['engines']))
    linuxpath, linux = latest('browser-*/results.json', lambda r: len(r['engines']) == 1 and r['engines'][0]['hostOS'] == 'linux')
    flowpath, flow = latest('workflows-*/results.json', lambda r: len(r['rows']) == 10 and all(t['status'] == 'PASS' for t in r['rows']))
    evalpath, evaluation = latest('workflows-*/results.json', lambda r: any(t['name'] == 'native-evaluation-tools-and-group-gradebook' and t['status'] == 'PASS' for t in r['rows']))
    unitpath, unit = latest('unit/*.json', lambda r: r.get('tests') == 60 and r['status'] == 'PASS')
    loadpath, load = latest('load-*.json', lambda r: r['status'] == 'PASS')
    renderpath, render = latest('render50-*.json', lambda r: r['status'] == 'PASS')
    context = read(OUT / 'context-final.json')
    links, youtube, candidate = (read(OUT / n) for n in ['links.json', 'youtube-link.json', 'candidate.json'])
    assert win['counts']['pass'] == 40 and win['counts']['fail'] == 0 and not win['issues']
    assert linux['counts']['pass'] == 12 and linux['counts']['fail'] == 0 and not linux['issues']
    assert unit['tests'] == 60 and unit['assertions'] == 252
    assert candidate['status'] == 'PASS' and candidate['source_matches_runtime']
    for component, key in [('local/asodeltolstaff', 'source_sha256'), ('theme/rsmax', 'theme_source_sha256')]:
        base = ROOT / 'plugins' / component
        current = {p.relative_to(base).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in base.rglob('*') if p.is_file()}
        assert current == candidate[key], 'Candidate evidence is stale: ' + component
    assert context['plugin_version'] == '2026102506' and context['theme_version'] == '2026102500'
    assert context['course_hidden'] and context['category_hidden']
    assert all(int(c['visible']) == 0 and int(c['categoryvisible']) == 0 for c in context['hidden_fixtures'])
    assert links['status'] == 'PASS' and len(links['rows']) == 98 and len(links['menu']) >= 3
    assert youtube['status'] == 'PASS' and youtube['currentTime'] > 0.5
    assert load['distinct_authenticated_users'] == render['distinctUsers'] == 50
    assert all(p['failed'] == 0 and p['maxSeconds'] < 60 for p in render['observations'])
    assert len(context['catalog_mapping']) == 50 and all(r['course_matches'] for r in context['catalog_mapping'])
    assert 'OK (2 tests, 5 assertions)' in (OUT / 'theme-favicon-final.txt').read_text(encoding='utf-8-sig')
    quality = (OUT / 'node-quality-merged.txt').read_text(encoding='utf-8-sig')
    assert re.search(r'fail 0', quality)
    nodecount = int(re.search(r'tests (\d+)', quality)[1])
    refs = {name: p.relative_to(ROOT).as_posix() for name, p in [('windows', winpath), ('linux', linuxpath),
             ('workflows', flowpath), ('evaluation', evalpath), ('phpunit', unitpath), ('http50', loadpath), ('render50', renderpath)]}
    browserrefs = [refs['windows'], refs['linux']]
    flowref = [refs['workflows']]
    # Each tuple is reviewed against the literal original criterion, not a completion heuristic.
    updates = {
      'REQ-01-01': (FULL, 'Curso sintético representativo, archivos completos y reproducción real.', 'Carga completa en navegador y archivo por SHA-256; reloj de video avanza.', 'PASS en el perfil local ensayado; no certifica capacidad del VPS.', 'Aceptación y nueva medición en el entorno institucional autorizado.', ['scripts/staging/close19/render50.mjs'], [refs['render50'], refs['http50']]),
      'REQ-01-02': (FULL, '50 identidades y contextos independientes; contenido representativo.', '550 peticiones HTTP y 400 comprobaciones con renderizado/descarga/reproducción.', 'PASS final: máximo de renderizado %.3f s. Intentos previos fallidos conservados.' % max(p['maxSeconds'] for p in render['observations']), 'Perfil probado caliente y controlado; repetir en infraestructura final y en arranque frío.', ['scripts/staging/close19/load.py', 'scripts/staging/close19/render50.mjs'], [refs['http50'], refs['render50']]),
      'REQ-02-01': (WAIT, 'Recorridos en Windows y navegador real en Linux.', 'Chromium/Edge/Firefox en Windows y WebKit en Linux.', 'Dos sistemas operativos acreditados; macOS no ejecutado.', 'QA debe ejecutar el recorrido en macOS; emulación no lo sustituye.', ['scripts/staging/close19/browser.mjs', 'scripts/staging/close19/linux.py'], browserrefs),
      'REQ-02-02': (FULL, 'Cuatro motores reales con sus dependencias; sincronización nativa de formularios.', 'Login, curso, unidad, actividad, notas, entrega y mensaje en 4 navegadores y 3 tamaños.', '52/52 comprobaciones funcionales finales PASS, combinando Windows y Linux.', 'Aceptación institucional; WebKit no equivale a Safari en un Mac físico.', ['scripts/staging/close19/browser.mjs', 'scripts/staging/close19/linux.py'], browserrefs),
      'REQ-02-03': (FULL, 'Escritorio 1440×900, tablet 768×1024 y móvil 390×844.', 'Sin scroll horizontal en rutas comprobadas; entrega persistida con archivo y mensaje enviado.', 'PASS en los 12 escenarios de navegador/tamaño; tamaños emulados.', 'Revisión institucional y dispositivos físicos para la aceptación final.', ['scripts/staging/close19/browser.mjs'], browserrefs),
      'REQ-02-04': (FULL, 'Favicon de respaldo nativo; espera de finalización del JS tras guardar.', 'Captura de consola, pageerror y respuestas HTTP; denegaciones esperadas separadas.', 'Cero errores inesperados en ambas ejecuciones finales. Fallos anteriores conservados.', 'Repetir con contenido definitivo; hallazgos de accesibilidad no se declaran subsanados.', ['plugins/theme/rsmax/classes/output/core_renderer.php', 'plugins/theme/rsmax/tests/favicon_test.php', 'scripts/staging/close19/browser.mjs'], browserrefs + ['build/staging/close19/theme-favicon-final.txt']),
      'REQ-03-01': (BLOCK, 'Monitor horario real con UTC, comprobación de login y cadena SHA-256; servicio Docker persistente.', 'Chequeos reales puntuales y primera observación del servicio; no se retrofechan datos.', 'PASS de observaciones existentes; periodo contractual NO acreditado.', 'Interventoría/operación deben fijar periodo, disponibilidad y continuidad del servicio; acumular la ventana completa.', ['scripts/staging/close19/monitor.py', 'scripts/staging/close19/monitor-service.py', 'infra/local/close19/compose.yml'], ['build/staging/close19/availability.jsonl', 'build/staging/close19/monitor/availability.jsonl']),
      'REQ-04-02': (FULL, 'MP4 H.264/AAC y MP3 propios; enlace al ejemplo oficial de YouTube.', 'Relojes de MP4/MP3 avanzan en los cuatro motores; enlace nativo abre YouTube y avanza %.3f s.' % youtube['currentTime'], 'PASS. El ensayo previo de iframe se conserva; el criterio solicita un enlace de YouTube.', 'Probar nuevamente los materiales académicos definitivos, sin atribuir licencia institucional al ejemplo.', ['scripts/staging/close19/seed.py', 'scripts/staging/close19/youtube-link.mjs'], browserrefs + ['build/staging/close19/youtube-link.json']),
      'REQ-08-01': (FULL, 'Verificador autenticado con lista segura de GET y navegación nativa.', '98 rutas locales, Más abre/cierra, menú móvil abre/cierra y enlace autorizado visible; YouTube externo abre/reproduce.', '98/98 rutas PASS; descargas íntegras; accesos no autorizados denegados.', 'Repetir tras la carga del contenido oficial; alcance actual: curso sintético y rutas ensayadas.', ['scripts/staging/close19/links.mjs'], ['build/staging/close19/links.json', 'build/staging/close19/youtube-link.json'] + flowref),
      'REQ-08-05': (BLOCK, 'Contraste del catálogo: cinco unidades y 50 plantillas en tipos nativos permitidos.', 'Verificación de cada código, tipo, sección y curso; auxiliares separados.', '50/50 coincidencias estructurales. No se inventa qué significan los cinco componentes.', 'Dirección académica debe identificar los cinco componentes del Anexo 14, resolver 50/80 y 4/9 y aportar contenido.', ['catalog/af8/catalog.json', 'scripts/staging/close19/fixture.php'], ['build/staging/close19/context-final.json']),
      'REQ-12-01': (FULL, 'Soporte privado con referencia persistente, historial y notificación nativa.', 'Solicitante crea caso, administrador responde/resuelve; tercero denegado; sin destinatario no muestra formulario.', 'PASS en UI, PHPUnit, privacidad y dos sesiones concurrentes: un caso y una notificación.', 'Designar responsable institucional; no implica SLA ni SMTP externo.', ['plugins/local/asodeltolstaff/classes/support.php', 'plugins/local/asodeltolstaff/support.php', 'plugins/local/asodeltolstaff/classes/privacy/provider.php', 'plugins/local/asodeltolstaff/classes/privacy/support_data.php'], flowref + [refs['phpunit']] + browserrefs),
      'REQ-12-02': (BLOCK, 'Seguimiento mínimo implementado: solicitud, respuesta, resolución e historial privado.', 'Caso sintético realmente respondido y resuelto; timestamps y notificación al solicitante.', 'Flujo PASS; cumplimiento del plazo NO acreditado porque no hay SLA aprobado.', 'Institución debe designar atención, canal/remitente oficial, SLA y evidencia del servicio; SMTP debe validarse por separado.', ['plugins/local/asodeltolstaff/db/install.xml', 'plugins/local/asodeltolstaff/db/upgrade.php', 'plugins/local/asodeltolstaff/classes/local/form/support_reply_form.php'], flowref + browserrefs),
      'REQ-13-02': (FULL, 'Campos nativos de curso y metadatos de roles, sin otorgar capacidades.', 'Guardar, recargar y consultar código, nombre, resumen, visibilidad, fechas, idioma y ambos identificadores; estudiante denegado.', 'PASS UI + API + consulta posterior. Curso permanece oculto.', 'Identificadores definitivos y aceptación institucional; los valores de QA no son roles oficiales.', ['plugins/local/asodeltolstaff/classes/metadata_fields.php', 'plugins/local/asodeltolstaff/tests/native_academic_fields_test.php'], flowref + ['build/staging/close19/context-final.json', refs['phpunit']]),
      'REQ-14-02': (FULL, 'Formulario nativo de grupo y campo administrativo de disponibilidad.', 'Nombre, numeración, descripción y estado se guardan/recargan; vínculo al curso consultado; estudiante denegado.', 'PASS; disponible/no disponible probado y disponible restaurado.', 'Los metadatos de disponibilidad no suspenden matrículas ni reemplazan restricciones académicas.', ['plugins/local/asodeltolstaff/classes/metadata_fields.php'], flowref + ['build/staging/close19/context-final.json', refs['phpunit']]),
      'REQ-14-03': (PART, 'Actividades, cuestionarios, foro, encuesta y libro nativo de notas; inventario técnico del grupo.', 'Beneficiario abre las cuatro herramientas; capacitador abre libro filtrado, otro grupo excluido y estudiante no abre grader.', 'PASS técnico; matriz académica de evaluación definitiva NO acreditada.', 'Dirección académica/tercero deben aportar matriz oficial, actividades e instrumentos y reglas aprobadas; no inventar pesos o notas.', ['scripts/staging/close19/workflows.mjs', 'plugins/local/asodeltolstaff/classes/group_publisher.php'], [refs['evaluation'], 'build/staging/close19/native-evaluation-matrix.json']),
      'REQ-19-02': (FULL, 'Anuncio y evento de cronograma mediante APIs nativas y formulario del capacitador.', 'Beneficiario del grupo ve anuncio y calendario; otro grupo no; 50 destinatarios únicos.', 'PASS UI y persistencia, sin fechas institucionales inventadas.', 'Publicar el cronograma institucional aprobado en su fase autorizada.', ['plugins/local/asodeltolstaff/communication.php', 'plugins/local/asodeltolstaff/classes/group_communication.php'], flowref),
      'REQ-19-03': (FULL, 'Condiciones hideIf admitidas por Moodle y práctica sin nota con retroalimentación textual nativa.', 'Capacitador sube MP4/PDF y crea actividad; beneficiario accede, otro grupo queda denegado.', 'PASS. Corregido fallo real Cannot load the grade item sin elegir notas oficiales.', 'Cargar material académico autorizado y definir sus reglas de evaluación.', ['plugins/local/asodeltolstaff/classes/local/form/publish_form.php', 'plugins/local/asodeltolstaff/classes/group_publisher.php', 'plugins/local/asodeltolstaff/tests/group_activity_notification_test.php'], flowref + [refs['phpunit']]),
      'REQ-19-05': (FULL, 'Apertura explícita de unidad para un grupo, preservando restricciones existentes.', 'Unidad abierta accesible al grupo; otro grupo denegado; condición previa de fecha conserva bloqueo.', 'PASS UI, JSON nativo de disponibilidad y PHPUnit.', 'Aplicar únicamente el calendario y las autorizaciones institucionales definitivas.', ['plugins/local/asodeltolstaff/classes/group_communication.php', 'plugins/local/asodeltolstaff/tests/section_release_test.php'], flowref + [refs['phpunit']]),
      'REQ-20-02': (FULL, 'Notificaciones internas nativas de anuncio y actividad del grupo.', 'Publicaciones reales por UI; 50 receptores únicos por evento; consulta de notificación del beneficiario.', 'PASS, sin destinatarios del otro grupo ni duplicados en cada publicación.', 'No equivale a entrega SMTP ni a una aprobación institucional.', ['plugins/local/asodeltolstaff/classes/group_communication.php', 'plugins/local/asodeltolstaff/classes/group_publisher.php'], flowref + [refs['phpunit']]),
    }
    assert len(updates) == 19
    tracking = copy.deepcopy(baseline)
    for row in tracking['rows']:
        if row['id'] not in updates:
            continue
        assert row['state'] == 'PENDIENTE'
        status, implementation, test, observed, dependency, files, evidence = updates[row['id']]
        row['historical_snapshot'] = copy.deepcopy(row)
        row.update(previous_state='PENDIENTE', technical_status=status,
                   state='VERIFICADO LOCAL' if status == FULL else 'BLOQUEADO' if status == BLOCK else 'PENDIENTE',
                   implementation=implementation, files=files, positive_test=test, observed=observed,
                   evidence=evidence, environment='http://127.0.0.1:8054 — Moodle 5.3+ (20261021), PHP 8.4.27, MySQL 8.4.13',
                   plugin_version='0.5.0 (2026102506)', theme_version='RSMax 1.2.4 (2026102500)',
                   production='NO ACCEDIDA NI MODIFICADA POR ESTA TAREA', dependency=dependency,
                   next_action=dependency, formal_acceptance='NO ACREDITADA', commit=git('rev-parse', '3326e77c2'))
    tracking['counts'] = dict(collections.Counter(r['state'] for r in tracking['rows']))
    technical = dict(collections.Counter(r['technical_status'] for r in tracking['rows'] if r['id'] in updates))
    assert sum(tracking['counts'].values()) == 82 and len({r['id'] for r in tracking['rows']}) == 82
    assert technical == {FULL: 14, WAIT: 1, BLOCK: 3, PART: 1}
    tracking.update(formal_acceptances=0, closure19={'generated_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'baseline_commit': BASE, 'technical_counts': technical, 'formal_acceptances': 0, 'production_actions': 0,
        'scope': '14 nuevos verificados en staging; 20 verificados históricos conservados, no recertificados hoy.'})
    dump(DOC / 'TRAZABILIDAD_82_2026-10-25.json', tracking)
    dump(OUT / 'tracking82.json', tracking)
    summary = {'environment': {k: context[k] for k in ['plugin_version','theme_version','moodle','php','mysql','hidden_fixtures']},
        'requirements82': tracking['counts'], 'closure19': technical, 'formal_acceptances': 0,
        'phpunit_distinct': {'tests': 62, 'assertions': 257, 'failures': 0, 'plugin': unit, 'theme': '2 tests, 5 assertions'},
        'node_tests': nodecount, 'browser_final': {'pass': 52, 'fail': 0, 'unexpected_errors': 0, 'engines': win['engines']+linux['engines'], 'evidence': browserrefs},
        'workflows': {'pass': 11, 'fail': 0, 'evidence': [refs['workflows'], refs['evaluation']]},
        'http50': {'requests': sum(p['requests'] for p in load['phases'].values()), 'status': load['status'], 'phases': load['phases']},
        'render50': {'checks': sum(len(p['rows']) for p in render['observations']), 'status': render['status'], 'phases': [{k:v for k,v in p.items() if k!='rows'} for p in render['observations']]},
        'links': {'checked': len(links['rows']), 'status': links['status'], 'menu': links['menu']}, 'youtube': youtube,
        'candidate': {'evidence': 'build/staging/close19/candidate.json', 'lint_checks': len(candidate['lint']), 'status': candidate['status']},
        'boundaries': ['No VPS access, deployment or formal acceptance.', 'Windows and Linux only; no macOS or physical device proof.',
                       'Earlier failed trials are retained, including resource contention and WebKit/Windows media decoding.',
                       'Accessibility scan found color-contrast and region violations; no WCAG certification.',
                       'Monitoring period, official evaluation matrix and five-component definition remain unaccredited.']}
    dump(PUBLIC / 'summary.json', summary)
    dump(PUBLIC / 'native-evaluation-matrix.json', read(OUT / 'native-evaluation-matrix.json'))
    dump(PUBLIC / 'candidate-hashes.json', candidate)
    history = []
    for pattern, key in [('unit/*.json','status'), ('load-*.json','status'), ('render50-*.json','status')]:
        for p in sorted(OUT.glob(pattern)):
            r = read(p); history.append({'path': p.relative_to(ROOT).as_posix(), 'status': r[key]})
    for p in sorted(OUT.glob('browser-*/results.json')):
        r = read(p); history.append({'path': p.relative_to(ROOT).as_posix(), 'counts': r['counts'], 'unexpected_issues': len(r['issues'])})
    for p in sorted(OUT.glob('workflows-*/results.json')):
        r = read(p); history.append({'path': p.relative_to(ROOT).as_posix(), 'pass': sum(t['status']=='PASS' for t in r['rows']), 'fail': sum(t['status']=='FAIL' for t in r['rows'])})
    dump(PUBLIC / 'execution-history.json', history)
    header = '# Matriz ASODELTOL AF8 — 82 criterios\n\n'
    header += 'Actualización del cierre técnico del 25 de octubre de 2026 (America/Bogota). Fuentes e IDs originales conservados. ASOCODEM-AF1 es otra institución; no acredita ASODELTOL.\n\n'
    header += '**34 verificados locales**: 20 históricos conservados y 14 nuevos demostrados en staging. **2 pendientes, 46 bloqueados, 0 aceptaciones oficiales.** Ninguno de estos estados certifica producción.\n\n'
    header += 'Detalle del lote: [CIERRE_19_2026-10-25.md](CIERRE_19_2026-10-25.md). Registro canónico: [TRAZABILIDAD_82_2026-10-25.json](TRAZABILIDAD_82_2026-10-25.json). Evidencia histórica conservada en cada registro y en Git; los antiguos archivos `build/staging` pertenecen al árbol original ASODELTOL.\n\n'
    header += 'Persisten las discrepancias de 50/80 ítems, cuatro/nueve recursos generales y la definición de cinco componentes por módulo. No se resuelven por inferencia.\n\n'
    header += '| ID | Criterio exacto de la preauditoría | Prioridad | Estado | Alcance verificado | Preparación técnica | Responsable pendiente | Archivos | Evidencia | Pendiente |\n|---|---|---|---|---|---|---|---|---|---|\n'
    states = {'VERIFICADO LOCAL':'IMPLEMENTADO+VERIFICADO','PENDIENTE':'PENDIENTE','BLOQUEADO':'BLOQUEADO','IMPLEMENTADO SIN VERIFICAR':'IMPLEMENTADO SIN VERIFICAR'}
    for r in tracking['rows']:
        current = r['id'] in updates
        values = [r['id'],r['criterion'],priorities[r['id']],states[r['state']],
                  r['observed'] if current else 'Histórico conservado; no recertificado hoy',
                  r.get('technical_status','Estado histórico conservado'),
                  'Ver dependencia específica' if current else 'Responsable histórico',r['files'],r['evidence'],r['next_action']]
        header += '| '+' | '.join(cell(v) for v in values)+' |\n'
    (DOC / 'MATRIZ_ASODELTOL_82.md').write_text(header, encoding='utf-8')
    closure = '# Cierre técnico de los 19 pendientes — 25 de octubre de 2026\n\n'
    closure += '**14 verificados integralmente en el perfil de staging, 1 con evidencia parcial, 1 pendiente de prueba adicional, 3 bloqueados por insumos/decisiones externas y 0 no conformes abiertos en el perfil final. Aceptaciones oficiales: 0.**\n\n'
    closure += 'Ambiente independiente: `127.0.0.1:8054`, Moodle 5.3+ (20261021), PHP 8.4.27, MySQL 8.4.13; plugin **0.5.0 / 2026102506**, RSMax **1.2.4 / 2026102500**. Cuatro cursos sintéticos y sus categorías permanecen ocultos. No se accedió al VPS ni se modificaron datos académicos reales desde esta tarea.\n\n'
    closure += '| ID | Estado anterior | Implementación realizada | Prueba ejecutada | Resultado | Estado final | Pendiente para cierre |\n|---|---|---|---|---|---|---|\n'
    for r in tracking['rows']:
        if r['id'] in updates:
            closure += '| '+' | '.join(cell(v) for v in [r['id'],'PENDIENTE',r['implementation'],r['positive_test'],r['observed'],r['technical_status'],r['dependency']])+' |\n'
    closure += '\nPruebas finales: **62 pruebas PHPUnit y 257 aserciones**, **%d pruebas Node**, **52 comprobaciones de navegador**, **11 flujos funcionales**, **98 rutas**, **550 peticiones HTTP de carga** y **400 comprobaciones renderizadas con 50 identidades**. Todos esos resultados finales son PASS dentro de su alcance. No sumar las repeticiones de suites como pruebas distintas.\n\n' % nodecount
    closure += 'Pruebas y huellas: [summary.json](evidencia-cierre19/summary.json), [historial completo de ejecuciones](evidencia-cierre19/execution-history.json), [índice de evidencia local](evidencia-cierre19/evidence-index.json), [candidato probado](evidencia-cierre19/candidate-hashes.json). Datos privados, claves y sesiones no se publican. El Excel actualizado está en `build/staging/close19/ASODELTOL_82_cierre19.xlsx`.\n\n'
    closure += 'Defectos corregidos: `hideIf notin` ocultaba campos de publicación; la práctica sin nota no tenía elemento nativo de retroalimentación y fallaba al abrirse; el favicon de respaldo de RSMax devolvía 404. Se implementó además seguimiento de soporte con referencia, respuesta, resolución, bloqueo concurrente y Privacy API. Los formularios y controles de acceso nativos se conservaron.\n\n'
    closure += 'Límites: WebKit/Windows no decodifica estos archivos incluso sin Moodle; WebKit/Linux sí los reproduce. Las primeras pruebas contenían fallos de sincronización del runner y un envío concurrente sobre la misma entrega sintética; se corrigieron usando identidades separadas y espera del JS nativo. Una carga renderizada anterior superó 60 s con contención local; la última prueba controlada completa aprobó, con máximo %.3f s. No acredita capacidad productiva, arranque frío universal ni dispositivos físicos. El ensayo inicial de iframe de YouTube no acreditó reproducción; el enlace nativo solicitado sí la acreditó posteriormente.\n\n' % max(p['maxSeconds'] for p in render['observations'])
    closure += 'El escaneo de accesibilidad encontró dos reglas con hallazgos (`color-contrast` y `region`); ejecutar el escaneo no significa aprobar WCAG. Se conservan en los artefactos. Behat global y macOS no se ejecutaron; las pruebas UI descritas son Playwright real. El monitor horario Docker está activo; la ventana contractual y continuidad 24/7 siguen sin acreditarse. El primer proceso Windows presentó una brecha superior a tres horas; se conservó su evidencia y se detuvo al sustituirlo por el servicio Docker.\n\n'
    closure += 'Publicación: cambios funcionales `3326e77c2` y `257899db1`; batería inicial `f86ebcb8c`; se integró `main` hasta `ced5aba9b`, preservando sus cambios ajenos. Durante la tarea, PR #14/#15 sustituyeron por otra misión el bloqueo histórico de despliegue. Esta entrega no revierte ese trabajo: la política actual clasifica ambos componentes como **critical / manualRequired=true**. No se fusiona automáticamente esta PR ni se ejecuta despliegue; la migración aditiva de dos tablas y su Privacy API requieren revisión y respaldo/recuperación coherentes antes de cualquier activación autorizada. No hacer downgrade directo ni borrar tablas para revertir.\n\n'
    closure += 'Solicitudes al tercero: entregar y cargar los materiales definitivos y licenciados, enlaces y accesibilidad; asociar cada recurso al catálogo y ejecutar nuevamente las pruebas tras la carga autorizada.\n\n'
    closure += 'Solicitudes a dirección académica: definición de los cinco componentes, aclaración de 50/80 y 4/9; matriz de evaluación e instrumentos, reglas/ponderaciones/notas, calendario y modalidad aprobados. No confundir los identificadores descriptivos probados con asignaciones oficiales de roles.\n\n'
    closure += 'Solicitudes a correo/SMTP: destinatario y remitente institucionales, responsable de atención y entrega SMTP verificable. Las notificaciones internas aprobadas no demuestran entrega por correo.\n\n'
    closure += 'Solicitudes a interventoría/operación: periodo de monitoreo, disponibilidad/SLA y evidencia del servicio de soporte; equipo macOS para la prueba faltante y aceptación formal de cada criterio. GitHub, staging y producción son estados separados.\n'
    (DOC / 'CIERRE_19_2026-10-25.md').write_text(closure, encoding='utf-8')
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book=Workbook(); sheet=book.active; sheet.title='ASODELTOL 82'
    sheet.append(['ID','Criterio original','Prioridad','Estado anterior','Estado actual','Estado técnico','Implementación','Prueba','Resultado','Evidencia','Pendiente','Aceptación oficial'])
    for r in tracking['rows']:
        sheet.append([r['id'],r['criterion'],priorities[r['id']],r.get('previous_state'),r['state'],r.get('technical_status','Histórico conservado'),r['implementation'],r['positive_test'],r['observed'],cell(r['evidence']),r['dependency'],'NO ACREDITADA'])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width=18 if column[0].column<2 else 45
        for c in column:
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(c.value,str): c.data_type='s'
    for c in sheet[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='174D40')
    for row in sheet.iter_rows(min_row=2): row[4].fill=PatternFill('solid',fgColor='C6EFCE' if row[4].value=='VERIFICADO LOCAL' else 'FFEB9C')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    book.save(OUT/'ASODELTOL_82_cierre19.xlsx')
    index = [{'path': p.relative_to(ROOT).as_posix(), 'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
             for p in sorted(OUT.rglob('*')) if p.is_file() and 'private' not in p.relative_to(OUT).parts]
    dump(PUBLIC / 'evidence-index.json', {'snapshot_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'note': 'Local artifacts only; private/ excluded. For append-only logs that grew, verify SHA-256 of the first bytes bytes recorded in each entry. Later observations are outside this snapshot.', 'files': index})
    print(json.dumps({'requirements82':tracking['counts'],'closure19':technical,'formal_acceptances':0},ensure_ascii=False))


if __name__=='__main__':
    main()
