"""Small, valid synthetic download fixtures; no institutional learning content."""
import datetime
import hashlib
import io
import json
import pathlib
import re
import zipfile
import struct
import zlib
import openpyxl
from pptx import Presentation
from pptx.util import Inches

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = ROOT / 'build/staging/moodle53/artifacts/anexo14-files'
OUT.mkdir(parents=True, exist_ok=True)
TEXT = 'ASODELTOL SYNTHETIC QA ONLY - NOT OFFICIAL MATERIAL'
FIXED = datetime.datetime(2026, 10, 24)


def package(entries):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in sorted(entries.items()):
            info = zipfile.ZipInfo(name, (2026, 10, 24, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content)
    return stream.getvalue()


def normalize(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
        if 'docProps/core.xml' in entries:
            entries['docProps/core.xml'] = re.sub(rb'(<dcterms:(?:created|modified)\b[^>]*>)[^<]+',
                rb'\g<1>2026-10-24T00:00:00Z', entries['docProps/core.xml'])
        return package(entries)


def pdf():
    stream = ('BT /F1 10 Tf 20 100 Td (' + TEXT + ') Tj ET').encode()
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 420 144] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    result = b'%PDF-1.4\n'
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result += str(index).encode() + b' 0 obj\n' + obj + b'\nendobj\n'
    start = len(result)
    result += b'xref\n0 6\n0000000000 65535 f \n'
    result += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    return result + f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n'.encode()


def png():
    def chunk(kind, content):
        return struct.pack('!I', len(content)) + kind + content + struct.pack('!I', zlib.crc32(kind + content) & 0xffffffff)
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!IIBBBBB', 1, 1, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(b'\x00\x30\x80\xc0')) + chunk(b'IEND', b''))


def main():
    word = package({
        '[Content_Types].xml': '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        '_rels/.rels': '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        'word/document.xml': '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>' + TEXT + '</w:t></w:r></w:p></w:body></w:document>',
    })
    workbook = openpyxl.Workbook()
    workbook.active.append([TEXT])
    workbook.properties.creator = 'ASODELTOL synthetic QA'
    workbook.properties.created = FIXED
    workbook.properties.modified = FIXED
    excel = io.BytesIO()
    workbook.save(excel)
    presentation = Presentation()
    presentation.core_properties.author = 'ASODELTOL synthetic QA'
    presentation.core_properties.created = FIXED
    presentation.core_properties.modified = FIXED
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = TEXT
    powerpoint = io.BytesIO()
    presentation.save(powerpoint)
    data = {'synthetic.docx': word, 'synthetic.xlsx': normalize(excel.getvalue()),
            'synthetic.pptx': normalize(powerpoint.getvalue()), 'synthetic.pdf': pdf(),
            'synthetic.png': png()}
    entries = []
    for filename, content in data.items():
        target = OUT / filename
        if target.exists():
            assert target.read_bytes() == content, 'Existing fixture preserved; content differs: ' + filename
        else:
            target.write_bytes(content)
        entries.append({'filename': filename, 'sha256': hashlib.sha256(content).hexdigest(), 'bytes': len(content)})
    (OUT / 'manifest.json').write_text(json.dumps({'synthetic': True, 'official_material': False, 'files': entries}, indent=2), encoding='utf-8')
    print(json.dumps({'synthetic_files': len(entries), 'official_material': False}))


if __name__ == '__main__':
    main()
