"""Rebuild the 82-row local compliance inventory from the ASODELTOL workbook.

Requires the already available openpyxl package. Does not change source books.
"""
import argparse
import collections
import hashlib
import json
import pathlib
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
STATES = ['IMPLEMENTADO+VERIFICADO', 'IMPLEMENTADO SIN VERIFICAR', 'PENDIENTE', 'BLOQUEADO']


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--sources', type=pathlib.Path, default=pathlib.Path.home() / 'Downloads')
    args = parser.parse_args()
    names = ['ASODELTOL_pre_auditoria_82_requisitos_2026-10-24.xlsx',
             '1-Verificacion_Anexo14_ASOCODEM_AF1_primera_revision (1).xlsx',
             'Estructura_AF8_ASODELTOL (1).docx']
    sources = []
    for name in names:
        path = args.sources / name
        if not path.is_file():
            raise SystemExit('Fuente requerida ausente: ' + name)
        sources.append({'name': name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    workbook = openpyxl.load_workbook(args.sources / names[0], data_only=True)
    values = list(workbook['Matriz 82'].values)
    assert len(values) == 83 and len({r[0] for r in values[1:]}) == 82
    # ASOCODEM is inspected for provenance, never used for ASODELTOL results.
    reference = openpyxl.load_workbook(args.sources / names[1], data_only=True)
    assert reference['Requisitos'].max_row == 83
    with zipfile.ZipFile(args.sources / names[2]) as doc:
        xml = ET.fromstring(doc.read('word/document.xml'))
        text = ' '.join(n.text for n in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if n.text)
    assert 'ASODELTOL' in text and 'AF8' in text
    rows = []
    hiddenpath = ROOT / 'build/staging/anexo14-qa/hidden-course.json'
    hidden = json.loads(hiddenpath.read_text(encoding='utf-8-sig')) if hiddenpath.exists() else {}
    qafiles = sorted((ROOT / 'build/staging/anexo14-qa').glob('qa_*.json'))
    latestqa = str(qafiles[-1].relative_to(ROOT)).replace('\\', '/') if qafiles else 'Sin QA HTTP'
    browserfiles = sorted((ROOT / 'build/staging/anexo14-browser').glob('*/summary.json'))
    latestbrowser = str(browserfiles[-1].relative_to(ROOT)).replace('\\', '/') if browserfiles else ''
    unitreports = sorted((ROOT / 'build/staging/anexo14-isolated').glob('*-summary.json'))
    unitreport = json.loads(unitreports[-1].read_text(encoding='utf-8')) if unitreports else {}
    unitxml = unitreports[-1].with_name(unitreports[-1].name.replace('-summary.json', '-junit.xml')) if unitreports else None
    cases = set()
    if unitreport.get('status') == 'PASS' and unitxml and unitxml.exists():
        cases = {c.attrib['name'] for c in ET.parse(unitxml).iter('testcase') if not any(c.find(tag) is not None for tag in ('error', 'failure', 'skipped'))}
    native = {
        'REQ-13-01': 'test_req_13_01_course_lifecycle',
        'REQ-14-01': 'test_req_14_01_group_lifecycle',
        'REQ-15-01': 'test_req_15_01_user_lifecycle',
        'REQ-16-01': 'test_req_16_01_02_individual_and_bulk_assignment',
        'REQ-16-02': 'test_req_16_01_02_individual_and_bulk_assignment',
        'REQ-17-02': 'test_req_17_02_custom_role_capabilities',
        'REQ-08-08': 'test_req_08_08_sequential_access',
        'REQ-19-04': 'test_req_19_04_grade_visible_to_learner',
        'REQ-20-04': 'test_req_20_04_21_03_bidirectional_messaging',
        'REQ-21-03': 'test_req_20_04_21_03_bidirectional_messaging',
    }
    reports = sorted((ROOT / 'build/staging/anexo14-reports').glob('*/summary.json'))
    reportresult = json.loads(reports[-1].read_text(encoding='utf-8')) if reports else {}
    reportchecks = {c['id']: c['status'] for c in reportresult.get('checks', [])}
    functional = sorted((ROOT / 'build/staging/anexo14-functional').glob('*/summary.json'))
    functionalresult = json.loads(functional[-1].read_text(encoding='utf-8')) if functional else {}
    functionalchecks = {c['id']: c['status'] for c in functionalresult.get('checks', [])}
    functionalmap = {
        'REQ-04-01': ['FILE-DOCX', 'FILE-XLSX', 'FILE-PPTX', 'FILE-PDF', 'FILE-PNG'],
        'REQ-14-04': ['GROUP-FORUM'],
        'REQ-19-01': ['TRAINER-GROUP-ISOLATION', 'GROUP-PAGE'],
        'REQ-19-06': ['PROGRESS-ROLE-SCOPE'],
        'REQ-20-03': ['LEARNER-GRADES-OWN-ONLY', 'PROGRESS-ROLE-SCOPE'],
        'REQ-21-02': ['AUDITOR-GRADES-READ-ONLY', 'PROGRESS-ROLE-SCOPE'],
    }
    goodhelp = []
    goodfiles = []
    goodadmin = []
    for path in browserfiles:
        result = json.loads(path.read_text(encoding='utf-8'))
        checks = {c['id']: c['status'] for c in result.get('checks', [])}
        if all(checks.get('HELP-' + size) == 'PASS' for size in ('desktop', 'tablet', 'mobile')) and checks.get('CONSOLE-NETWORK') == 'PASS':
            goodhelp.append(path)
        if all(checks.get('PRIVATE-FILES-' + size) == 'PASS' for size in ('desktop', 'mobile')) and checks.get('CONSOLE-NETWORK') == 'PASS':
            goodfiles.append(path)
        if all(checks.get('STAFF-' + size) == 'PASS' for size in ('desktop', 'tablet', 'mobile')) and checks.get('CONSOLE-NETWORK') == 'PASS':
            goodadmin.append(path)
    if goodadmin:
        latestbrowser = goodadmin[-1].relative_to(ROOT).as_posix()
    # Only these criteria have a concrete native/plugin implementation covering
    # the requested function. Partial structural support remains PENDIENTE.
    implemented = {'REQ-04-01', 'REQ-13-01', 'REQ-14-01', 'REQ-15-01', 'REQ-16-01',
                   'REQ-17-02', 'REQ-19-01', 'REQ-19-04', 'REQ-20-01', 'REQ-20-04', 'REQ-21-03'}
    for number, row in enumerate(values[1:], 2):
        rid, item, section, mandatory, actor, title, criterion, oldstate, priority, files, pending = row
        state = 'PENDIENTE'
        evidence = ''
        scope = 'Sin verificación integral'
        if oldstate in ('Dependencia pendiente', 'Aplicabilidad por definir'):
            state = 'BLOQUEADO'
        elif rid in implemented:
            state = 'IMPLEMENTADO SIN VERIFICAR'
        if rid == 'REQ-18-01' and hidden.get('status') == 'PASS':
            state = 'IMPLEMENTADO+VERIFICADO'
            scope = 'HTTP local; beneficiario sintético autenticado; curso y categoría ocultos'
            files += '; scripts/staging/check-hidden-course.ps1'
            evidence = 'build/staging/anexo14-qa/hidden-course.json'
            pending = 'Aceptación institucional y verificación en entorno final; la prueba local no acredita producción.'
        if rid.startswith('REQ-02-') and latestbrowser:
            files += '; scripts/staging/check-staff-browser.mjs'
            evidence = latestbrowser
            pending = 'Edge/Chromium local: login y listado staff, escritorio/tablet/móvil. Faltan demás SO/navegadores, recorrido completo del curso, entrega y mensajería; no se cierra el criterio.'
        if rid in ('REQ-19-01', 'REQ-23-03', 'REQ-24-03', 'REQ-24-04'):
            files += '; plugins/local/asodeltolstaff/classes/manager.php; plugins/local/asodeltolstaff/export.php; plugins/local/asodeltolstaff/tests/manager_test.php'
            evidence = 'build/staging/moodle53/mission-11-evidence.json; build/staging/anexo14-qa/m11-service-recheck.json; ' + latestqa
            if rid == 'REQ-19-01':
                pending = 'Evidencia parcial M11: aislamiento del plugin; falta demostrar el recorrido completo del curso por perfil.'
            elif rid == 'REQ-24-04':
                pending = 'Export sintético ampliado con inicio/fin programado y reportado UTC, reconciliado CSV/XLSX. No son login/cierre ni inicio/fin de acción o disponibilidad de dudas. Falta completar reporte contractual.'
                evidence += '; build/staging/anexo14-qa/export-content.json'
            elif rid in ('REQ-23-03', 'REQ-24-03'):
                pending = 'Plugin y export sintético operativos; faltan fuente de asistencia efectiva, actividades realizadas y formato institucional. No usar logs de acceso como asistencia.'
        if rid in native and native[rid] in cases:
            state = 'IMPLEMENTADO+VERIFICADO'
            scope = 'API nativa Moodle 5.3; base/cuenta/dataroot aislados; fixtures sintéticos; no interfaz ni roles institucionales definitivos'
            evidence = unitxml.relative_to(ROOT).as_posix() + '#' + native[rid]
            files += '; plugins/local/asodeltolstaff/tests/anexo14_native_test.php'
            pending = 'Aceptación institucional y recorrido final con configuración/cuentas autorizadas; no se acredita producción.'
        if rid == 'REQ-08-02' and goodhelp:
            state = 'IMPLEMENTADO+VERIFICADO'
            scope = 'Beneficiario existente en HTTP local; enlace visible y clic real; escritorio/tablet/móvil'
            evidence = '; '.join(p.relative_to(ROOT).as_posix() for p in goodhelp[-2:])
            files += '; plugins/local/asodeltolstaff/help.php; plugins/local/asodeltolstaff/classes/hook_callbacks.php; plugins/local/asodeltolstaff/db/hooks.php'
            pending = 'Validar en la instalación final tras el despliegue autorizado; aceptación oficial pendiente.'
        if rid == 'REQ-20-01' and goodfiles:
            state = 'IMPLEMENTADO+VERIFICADO'
            scope = 'Beneficiario sintético; selector nativo, guardado y recarga; descarga por clic y SHA256 en perfiles escritorio/celular emulados; no hardware móvil'
            evidence = goodfiles[-1].relative_to(ROOT).as_posix()
            files += '; scripts/staging/check-staff-browser.mjs; Moodle core_user/form/private_files'
            pending = 'Aceptación institucional y comprobación en entorno final; no acredita entrega académica ni hardware móvil real.'
        if rid == 'REQ-24-04' and reportchecks.get('ACCESS') == 'PASS' and 'test_actual_sources_and_missing_logout_remain_distinct' in cases:
            state = 'IMPLEMENTADO+VERIFICADO'
            scope = 'Campos y fuentes nativas probados con fixtures completos y descarga HTTP real; logs independientes, no tiempo trabajado'
            files += '; plugins/local/asodeltolstaff/classes/access_report.php; plugins/local/asodeltolstaff/classes/office_hours.php'
            evidence = reports[-1].relative_to(ROOT).as_posix() + '; ' + unitxml.relative_to(ROOT).as_posix()
            pending = 'Configurar fechas y horas institucionales; aceptar interpretación del reporte. No inferir duración/continuidad con login/logout ausentes.'
        if (rid in functionalmap and functionalresult.get('status') == 'PASS'
                and functionalresult.get('final_hidden') is True
                and all(functionalchecks.get(check) == 'PASS' for check in functionalmap[rid])):
            state = 'IMPLEMENTADO+VERIFICADO'
            scope = 'HTTP local con perfiles sintéticos existentes; archivos por SHA256, notas/avance nativos y denegación entre grupos; no producción'
            evidence = functional[-1].relative_to(ROOT).as_posix()
            files += '; scripts/staging/check-native-course.py; scripts/staging/anexo14-functional-context.php'
            pending = 'Aceptación institucional y configuración final; fixtures sintéticos no son contenido ni calificaciones oficiales.'
        readiness = 'PROBADO LOCAL' if state == 'IMPLEMENTADO+VERIFICADO' else 'Revisar pendiente del criterio'
        owner = 'Equipo técnico / QA'
        if rid.startswith(('REQ-06-', 'REQ-07-')) or rid in ('REQ-08-03', 'REQ-08-04', 'REQ-08-05', 'REQ-08-06'):
            owner = 'Tercero de contenidos + responsable académico'
            readiness = 'Catálogo y plantilla de 54 componentes disponibles; carga y revisión reales pendientes'
        if rid == 'REQ-08-09':
            owner = 'Responsable académico / capacitador'
            readiness = 'Formulario y calendario nativo probados; pendiente horario oficial. La cita sintética no acredita el horario institucional.'
            files += '; plugins/local/asodeltolstaff/officehours.php; plugins/local/asodeltolstaff/classes/office_hours.php'
            evidence = 'build/staging/anexo14-reports/office-hours-http.json'
        if rid == 'REQ-19-03':
            readiness = 'Publicador propio de páginas/URL/archivos probado con capacidad dedicada sin concesiones en staging; faltan materiales y actividades oficiales'
            files += '; plugins/local/asodeltolstaff/classes/group_publisher.php; plugins/local/asodeltolstaff/publish.php'
            evidence = unitxml.relative_to(ROOT).as_posix() if unitxml else ''
        if rid in ('REQ-24-01', 'REQ-24-02') and reportchecks:
            readiness = 'Reporte base CSV/XLSX funcional y probado; diccionario Anexo12/regla oficial de culminación pendientes'
            owner = 'Institución / Interventoría para definición; equipo técnico para adaptación'
            files += '; plugins/local/asodeltolstaff/classes/report_manager.php; plugins/local/asodeltolstaff/reports.php'
            evidence = reports[-1].relative_to(ROOT).as_posix()
        if rid.startswith(('REQ-22-', 'REQ-15-02', 'REQ-16-03', 'REQ-17-01', 'REQ-20-05', 'REQ-21-01')):
            owner = 'Institución; roles y datos definitivos reservados a última fase'
        rows.append(dict(id=rid, item=item, section=section, requirement=title, criterion=criterion,
                         priority=priority, status=state, actor=actor, mandatory=mandatory,
                         verification_scope=scope, technical_readiness=readiness, pending_owner=owner,
                         files=files, evidence=evidence or 'Sin evidencia integral del criterio; ver soporte de archivos.',
                         pending=pending, official_acceptance='NO ACREDITADA', source_row=number,
                         source_status=oldstate))
    counts = {state: sum(r['status'] == state for r in rows) for state in STATES}
    out = ROOT / 'build/staging/anexo14-compliance'
    out.mkdir(parents=True, exist_ok=True)
    payload = {'scope': 'ASODELTOL AF8, staging local; estados del criterio completo',
               'sources': sources, 'counts': counts, 'official_acceptance': 'NO ACREDITADA', 'requirements': rows}
    (out / 'matrix-82.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'ASODELTOL 82'
    keys = list(rows[0])
    sheet.append(keys)
    for row in rows:
        sheet.append([row[k] for k in keys])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 28
    for row in sheet:
        for cell in row:
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
    summary = book.create_sheet('Resumen y fuentes')
    summary.append(['Estado', 'Cantidad'])
    for state, count in counts.items():
        summary.append([state, count])
    summary.append(['Aceptación oficial', 'NO ACREDITADA'])
    for source in sources:
        summary.append([source['name'], source['sha256']])
    book.save(out / 'ASODELTOL-matriz-82.xlsx')
    saved = openpyxl.load_workbook(out / 'ASODELTOL-matriz-82.xlsx', data_only=True)
    assert saved.worksheets[0].max_row == 83
    assert sum(counts.values()) == 82
    lines = ['# Matriz ASODELTOL AF8 — 82 criterios', '',
             'Fuente: libro de preauditoría ASODELTOL, hoja `Matriz 82`. ASOCODEM-AF1 es otra institución y solo referencia; no se trasladan sus resultados.', '',
             'Los estados corresponden a la función demostrada en local, con su alcance explícito. Una prueba API aislada no acredita interfaz, roles definitivos ni producción. Evidencia parcial no constituye cumplimiento. Aceptación oficial: **NO ACREDITADA** en los 82 criterios.', '',
             '| Estado | Total |', '|---|---:|']
    lines += [f'| {s} | {n} |' for s, n in counts.items()]
    lines += ['', 'Matriz editable y JSON: `build/staging/anexo14-compliance/`. Regenerar con `python scripts/anexo14/build-compliance-matrix.py`.', '',
              'Discrepancias DOCX conservadas: tabla inicial con totales 80/50; instrucciones de nueve ítems generales frente a cuatro enumerados; referencias a unidad 1 copiadas en otras unidades. Requieren aclaración académica; no se corrigen reglas oficiales por inferencia.', '',
              '| ID | Criterio exacto de la preauditoría | Prioridad | Estado | Alcance verificado | Preparación técnica | Responsable pendiente | Archivos | Evidencia | Pendiente |', '|---|---|---|---|---|---|---|---|---|---|']
    for row in rows:
        lines.append('| ' + ' | '.join(str(row[k]).replace('|', '\\|').replace('\n', ' ') for k in
                                      ('id', 'criterion', 'priority', 'status', 'verification_scope', 'technical_readiness', 'pending_owner', 'files', 'evidence', 'pending')) + ' |')
    lines += ['', '## Trazabilidad de fuentes', '']
    lines += [f"- `{s['name']}` — SHA256 `{s['sha256']}`" for s in sources]
    (ROOT / 'docs/anexo14/MATRIZ_ASODELTOL_82.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(json.dumps({'counts': counts, 'rows_verified': 82, 'xlsx_reopened': True}, ensure_ascii=False))


if __name__ == '__main__':
    main()
